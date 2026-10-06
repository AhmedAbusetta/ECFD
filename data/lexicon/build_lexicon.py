# -*- coding: utf-8 -*-
"""Build the team lexicon workbook + CSV from lexicon_seed.py.

  python data/lexicon/build_lexicon.py

Outputs (next to this file):
  ECFD_Egyptian_Call_Lexicon.xlsx  - upload to Google Sheets / share with the team
  egyptian_call_lexicon.csv        - machine-readable copy for the NLP matcher and the scorer
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from lexicon_seed import ROWS

HERE = Path(__file__).parent
XLSX = HERE / "ECFD_Egyptian_Call_Lexicon.xlsx"
CSV = HERE / "egyptian_call_lexicon.csv"

HEADERS = ["ID", "Word (as spoken)", "Other spellings (| separated)", "Meaning (English)", "Category",
           "Concept", "Risk role", "Said by", "Example sentence (Egyptian)", "Status", "Added by", "Notes"]
WIDTHS = [7, 20, 38, 28, 26, 22, 11, 10, 42, 18, 12, 40]
ARABIC_COLS = {2, 3, 9}  # right-aligned

CATEGORIES = sorted({r[3] for r in ROWS})
ROLES = ["HARD", "SCORE", "CONTEXT", "NEUTRAL"]
SAID_BY = ["Caller", "Employee", "Both"]
STATUSES = ["Claude draft", "Verified", "Needs discussion", "Rejected"]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
ROLE_FILL = {"HARD": "F4CCCC", "SCORE": "FCE5CD", "CONTEXT": "D9EAD3", "NEUTRAL": "FFFFFF"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def rows_with_ids():
    for i, (word, variants, meaning, cat, concept, role, said_by, example, note) in enumerate(ROWS, 1):
        yield [f"L{i:03d}", word, variants, meaning, cat, concept, role, said_by, example,
               "Claude draft", "Claude", note]


def build_readme(ws):
    lines = [
        ("ECFD Egyptian Call Lexicon", True),
        ("The words Egyptians actually use on the phone: Arabized English verbs (اتشيك، كنسل)، slang, request verbs, secret objects, scam phrases.", False),
        ("", False),
        ("What it is used for", True),
        ("1. NLP detection: one entry per stem; the matcher strips Arabic prefixes/suffixes (ه، ب، ات … ها، لي، هولي) so كنسل also catches هكنسلهالك.", False),
        ("2. ASR testing: every Example sentence is a line to record (Test 1 / bake-off).", False),
        ("3. ASR fine-tuning: the example sentences + recordings become training data.", False),
        ("", False),
        ("How to add or check words (yellow cells are for you)", True),
        ("• Add a new row at the bottom of the Lexicon tab. Fill Word, Other spellings, Meaning, Category, Risk role, Said by, Example sentence, Added by.", False),
        ("• Every row starts as 'Claude draft'. A native speaker changes Status to 'Verified' (or 'Rejected' / 'Needs discussion').", False),
        ("• Other spellings: every way the ASR might write it - Arabic script, Latin, with/without ال. Separate with |", False),
        ("• Example sentence: how a real person would say it on a call - not formal Arabic.", False),
        ("• Rows marked 'verify' in Notes (mostly money slang) are ones Claude is unsure about.", False),
        ("", False),
        ("Risk roles", True),
        ("HARD     part of a hard signal: a request verb (تديني، ابعتلي، قولّي) + a secret object (الكود، الcvv، الباسورد). Only HARD combos raise critical alerts.", False),
        ("SCORE    adds to the risk score only (pressure, pretexts, payment talk).", False),
        ("CONTEXT  context only - e.g. claiming to be from the bank is not proof of fraud (ADR-0004).", False),
        ("NEUTRAL  normal Egyptian speech; listed so the ASR is tested on it.", False),
        ("", False),
        ("Source of truth: data/lexicon/lexicon_seed.py in the repo. After editing the shared sheet, export it and the team lead merges the changes.", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=14 if r == 1 else 11, color="1F3864" if bold else "000000")
    ws.column_dimensions["A"].width = 150
    legend = ws.cell(row=len(lines) + 2, column=1, value="Legend: yellow cells = to fill in / verify")
    legend.fill = INPUT_FILL
    legend.font = Font(name=FONT, italic=True)


def build_lexicon(ws, data):
    ws.append(HEADERS)
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]
    for row in data:
        ws.append(row)
        r = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            c = ws.cell(row=r, column=col)
            c.font = Font(name=FONT, size=11)
            c.border = BORDER
            c.alignment = Alignment(horizontal="right" if col in ARABIC_COLS else "left",
                                    vertical="top", wrap_text=True, readingOrder=2 if col in ARABIC_COLS else 0)
        ws.cell(row=r, column=7).fill = PatternFill("solid", fgColor=ROLE_FILL[row[6]])
        ws.cell(row=r, column=10).fill = INPUT_FILL
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    last = 1000  # leave room for the team to add rows
    for col, options in ((5, CATEGORIES), (7, ROLES), (8, SAID_BY), (10, STATUSES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{last}")
        ws.add_data_validation(dv)


def build_summary(ws, n_rows):
    ws["A1"] = "Lexicon summary (live counts from the Lexicon tab)"
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color="1F3864")
    rng = lambda col: f"Lexicon!${col}$2:${col}$1000"

    ws["A3"], ws["B3"], ws["C3"] = "Category", "Words", "Verified"
    r = 4
    for cat in CATEGORIES:
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=f'=COUNTIF({rng("E")},A{r})')
        ws.cell(row=r, column=3, value=f'=COUNTIFS({rng("E")},A{r},{rng("J")},"Verified")')
        r += 1
    ws.cell(row=r, column=1, value="Total")
    ws.cell(row=r, column=2, value=f"=SUM(B4:B{r - 1})")
    ws.cell(row=r, column=3, value=f"=SUM(C4:C{r - 1})")
    total_row = r

    r += 2
    ws.cell(row=r, column=1, value="Risk role")
    ws.cell(row=r, column=2, value="Words")
    role_head = r
    for role in ROLES:
        r += 1
        ws.cell(row=r, column=1, value=role)
        ws.cell(row=r, column=2, value=f'=COUNTIF({rng("G")},A{r})')

    r += 2
    ws.cell(row=r, column=1, value="Verified so far")
    ws.cell(row=r, column=2, value=f"=IFERROR(C{total_row}/B{total_row},0)")
    ws.cell(row=r, column=2).number_format = "0%"

    for row in ws.iter_rows(min_row=3, max_row=r):
        for c in row:
            c.font = Font(name=FONT, bold=c.row in (3, total_row, role_head))
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12


def build_guide(ws):
    rules = [
        ("Transcription guideline (use the same rules for every recording)", ""),
        ("Rule", "Example"),
        ("Write Egyptian as spoken - never convert to formal Arabic", "عايز (not أريد)، دلوقتي (not الآن)"),
        ("English acronyms and brand names in Latin letters", "OTP, CVV, PIN, IT, AnyDesk, InstaPay"),
        ("ال attached to an English word: keep it attached, as spoken", "الcvv، الotp"),
        ("Arabized English verbs in Arabic letters, as pronounced", "اتشيك، هكنسلهالك، فورودهولي"),
        ("English nouns Egyptians use every day in Arabic letters", "الباسورد، اللينك، الأبلكيشن، الأكونت"),
        ("Digits spoken one by one: write each as a digit", "صفر واحد خمسة -> 0 1 5"),
        ("Keep Egyptian ج/ق as heard", "جالك (not قالك)"),
        ("Don't add punctuation you didn't hear; no diacritics", ""),
        ("Unclear word: write [?] - never guess", "ابعتلي [?] بسرعة"),
    ]
    for r, (a, b) in enumerate(rules, 1):
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b)
        for col in (1, 2):
            c = ws.cell(row=r, column=col)
            c.font = Font(name=FONT, bold=r <= 2, size=13 if r == 1 else 11, color="1F3864" if r == 1 else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right" if col == 2 else "left")
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 45


def main():
    data = list(rows_with_ids())
    wb = Workbook()
    build_readme(wb.active)
    wb.active.title = "How to use"
    build_lexicon(wb.create_sheet("Lexicon"), data)
    build_summary(wb.create_sheet("Summary"), len(data))
    build_guide(wb.create_sheet("Transcription guide"))
    wb.active = 1
    wb.save(XLSX)

    with CSV.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(data)
    print(f"{len(data)} entries -> {XLSX.name}, {CSV.name}")


if __name__ == "__main__":
    main()
